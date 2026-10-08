"""Immutable file receipt and resumable text extraction, outside write transactions."""
from __future__ import annotations
import hashlib
import json
import mimetypes
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import uuid
from . import core


def _safe_name(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).rstrip('. ')[:140] or 'file'


def save(root, path, *, title=None, document_id=None, source_id=None, source_url=None,
         authority='user', evidence_kind='user_provided', domain='general', topics=None,
         published_at=None, expected_revision=None, parse=True):
    base, source = core.root_path(root), Path(path).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    digest = core.sha256(source)
    if document_id is None:
        document_id = ('document:' + uuid.uuid5(uuid.NAMESPACE_URL, str(source_id or authority) + ':' + source_url).hex) if source_url else core.uid('document')
    version_id = 'version:' + uuid.uuid5(uuid.NAMESPACE_URL, document_id + ':' + digest).hex
    target = base / 'data/originals' / _safe_name(document_id) / _safe_name(version_id) / _safe_name(source.name)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        if core.sha256(target) != digest:
            raise RuntimeError('Immutable original at destination differs; refusing overwrite')
    else:
        temp = target.with_name('.' + target.name + '.' + uuid.uuid4().hex + '.partial')
        try:
            shutil.copy2(source, temp)
            if core.sha256(temp) != digest or core.sha256(source) != digest:
                raise RuntimeError('Source changed during receipt; retry from stable input')
            with temp.open('r+b') as stream:
                os.fsync(stream.fileno())
            # A concurrent identical receipt has the same bytes. No mutable user file is replaced.
            os.replace(temp, target)
        finally:
            temp.unlink(missing_ok=True)
    rel = target.relative_to(base).as_posix()
    con = core.connect(base)
    try:
        def write(c):
            existing_version = c.execute('SELECT * FROM versions WHERE version_id=?', (version_id,)).fetchone()
            if existing_version:
                if not existing_version['is_current']:
                    revision = core._bump(c, document_id, expected_revision)
                    old_document = dict(c.execute('SELECT * FROM documents WHERE document_id=?', (document_id,)).fetchone())
                    core._change(c, document_id, revision, old_document, {'current_version_id': version_id}, source_ref=source_url or rel, reason='Previously saved content observed again; reused original becomes current')
                    c.execute('UPDATE versions SET is_current=0 WHERE document_id=?', (document_id,))
                    c.execute('UPDATE versions SET is_current=1 WHERE version_id=?', (version_id,))
                    c.execute('UPDATE documents SET source_path=?,last_seen_at=?,source_updated_at=? WHERE document_id=?', (existing_version['raw_path'], core.now(), core.now(), document_id))
                    c.execute('UPDATE locations SET is_current=0 WHERE item_id=? AND role=?', (document_id, 'original'))
                    changed = c.execute('UPDATE locations SET is_current=1 WHERE item_id=? AND path=? AND role=?', (document_id,existing_version['raw_path'],'original')).rowcount
                    if not changed:
                        c.execute('INSERT INTO locations VALUES(?,?,?,?,?,?,?,?)', (core.uid('location'),document_id,existing_version['raw_path'],'original',0,1,core.now(),core.json_text({'version_id':version_id})))
                    from .search import reindex_item
                    reindex_item(c, document_id)
                return False
            stamp = core.now()
            sid = source_id or ('v2:receipt:' + authority)
            found_source = c.execute('SELECT authority FROM sources WHERE source_id=?', (sid,)).fetchone()
            if found_source and found_source['authority'] != authority:
                raise ValueError('Source authority mismatch; use registered source authority explicitly')
            if not found_source:
                c.execute('INSERT INTO sources(source_id,name,source_kind,base_url,authority,access_mode,scope,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)',
                          (sid, sid, 'manual_receipt', source_url, authority, 'local', 'Explicitly received files; not a coverage claim', stamp, stamp))
            old = c.execute('SELECT * FROM documents WHERE document_id=?', (document_id,)).fetchone()
            if old:
                revision = core._bump(c, document_id, expected_revision)
                core._change(c, document_id, revision, dict(old), {'new_version_id': version_id, 'sha256': digest}, source_ref=source_url or rel, reason='New immutable document version')
                c.execute('UPDATE versions SET is_current=0 WHERE document_id=?', (document_id,))
                c.execute('UPDATE search_entries SET is_current=0 WHERE item_id=? AND version_id!=?', (document_id, version_id))
                c.execute('UPDATE documents SET title=?,source_path=?,last_seen_at=?,source_updated_at=? WHERE document_id=?', (title or old['title'], rel, stamp, stamp, document_id))
                c.execute('UPDATE items SET title=? WHERE item_id=?', (title or old['title'], document_id))
            else:
                core.new_item(c, 'document', title or source.name, domain=domain, authority=authority, evidence_kind=evidence_kind, item_id=document_id)
                c.execute('''INSERT INTO documents(document_id,source_id,document_kind,title,canonical_url,source_path,published_at,first_seen_at,last_seen_at,evidence_kind)
                          VALUES(?,?,?,?,?,?,?,?,?,?)''', (document_id, sid, 'file', title or source.name, source_url, rel, published_at, stamp, stamp, evidence_kind))
            c.execute('INSERT INTO versions(version_id,document_id,content_sha256,raw_path,media_type,captured_at,byte_size,is_current) VALUES(?,?,?,?,?,?,?,1)',
                      (version_id, document_id, digest, rel, mimetypes.guess_type(source.name)[0] or 'application/octet-stream', stamp, target.stat().st_size))
            c.execute('INSERT INTO document_processing VALUES(?,?,?,?,?)', (version_id, 'received', None, None, stamp))
            c.execute('UPDATE locations SET is_current=0 WHERE item_id=?', (document_id,))
            c.execute('INSERT INTO locations VALUES(?,?,?,?,?,?,?,?)', (core.uid('location'), document_id, rel, 'original', 0, 1, stamp, core.json_text({'version_id': version_id})))
            for topic in topics or []:
                if not c.execute('SELECT 1 FROM topics WHERE topic_id=?', (topic,)).fetchone():
                    c.execute('INSERT INTO topics(topic_id,name,created_at,updated_at) VALUES(?,?,?,?)', (topic, topic, stamp, stamp))
                c.execute('''INSERT OR IGNORE INTO document_topics(document_id,topic_id,topic_role,assignment_method,confidence,created_at,updated_at)
                              VALUES(?,?,?,?,?,?,?)''', (document_id, topic, 'related', 'explicit', 1, stamp, stamp))
            current_topics = [r[0] for r in c.execute('SELECT topic_id FROM document_topics WHERE document_id=?', (document_id,))]
            core.index_entry(c, document_id, '\n'.join([title or source.name, source_url or '', source.name]), entry_key='legacy:version:' + version_id, source_ref=rel, version_id=version_id, topics=current_topics)
            return True
        created = core.transaction(con, write)
    finally:
        con.close()
    check = core.connect(base, readonly=True)
    try:
        saved_version = check.execute('SELECT raw_path FROM versions WHERE version_id=?', (version_id,)).fetchone()
        saved_processing = check.execute('SELECT state FROM document_processing WHERE version_id=?', (version_id,)).fetchone()
    finally:
        check.close()
    result = {'item_id': document_id, 'version_id': version_id, 'path': saved_version['raw_path'], 'sha256': digest, 'new_version': created, 'status': saved_processing['state'] if saved_processing else 'received'}
    if parse:
        result.update(parse_version(base, version_id))
    return result


def extract(path):
    suffix = path.suffix.lower()
    chunks, warnings = [], []
    if suffix in {'.doc', '.xls'}:
        office = shutil.which('soffice') or shutil.which('soffice.com')
        if not office:
            return [], ['LibreOffice unavailable; original retained without extracted text']
        target_suffix = '.docx' if suffix == '.doc' else '.xlsx'
        with tempfile.TemporaryDirectory(prefix='pis-office-') as temporary:
            directory = Path(temporary)
            profile = (directory / 'profile').as_uri()
            output = directory / (path.stem + target_suffix)
            command = [office, f'-env:UserInstallation={profile}', '--headless',
                       '--convert-to', target_suffix[1:], '--outdir', str(directory), str(path)]
            result = subprocess.run(command, capture_output=True, timeout=90)
            if result.returncode or not output.is_file():
                detail = (result.stderr or result.stdout or b'conversion produced no file').decode('utf-8', 'replace').strip()
                raise ValueError(f'LibreOffice conversion failed: {detail[:300]}')
            return extract(output)
    if suffix in {'.txt', '.md', '.json', '.jsonl', '.csv', '.tsv', '.tex', '.py', '.ps1', '.log', '.yaml', '.yml', '.toml', '.ini'}:
        raw = path.read_bytes()
        encodings = ('utf-16',) if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'gb18030')
        for encoding in encodings:
            try:
                text = raw.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError('Text encoding unsupported')
        # Preserve authorship for exported ChatGPT conversations.
        matches = list(re.finditer(r'^##\s+\d+\.\s+(你|ChatGPT)\s*$', text, re.M))
        if matches:
            for index, match in enumerate(matches):
                end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
                role = 'user' if match.group(1) == '你' else 'assistant'
                chunks.append((match.group(0), text[match.end():end].strip(), role, 'user_statement' if role == 'user' else 'assistant_analysis'))
        else:
            lines = text.splitlines()
            for start in range(0, len(lines), 80):
                chunks.append((f'lines {start+1}-{min(start+80,len(lines))}', '\n'.join(lines[start:start+80]), '', ''))
    elif suffix in {'.html', '.htm'}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(path.read_bytes(), 'html.parser')
        for tag in soup(['script', 'style', 'nav']):
            tag.decompose()
        body = (soup.select_one('.post-entry') or soup.select_one('.v_news_content')
                or soup.select_one('#vsb_content') or soup.select_one('.wp_articlecontent'))
        if body:
            content = body.get_text('\n', strip=True)
            embedded = [tag.get('src') or tag.get('data') or tag.get('pdfsrc')
                        for tag in body.select('img[src],iframe[src],embed[src],object[data],[pdfsrc]')]
            if len(content) < 20 and embedded:
                warnings.append('HTML body depends on linked PDF/images; retrieve and extract those assets before completion')
            if not content and not embedded:
                warnings.append('HTML article container is empty; title metadata is not full text')
            if not content:
                content = '\n'.join([(soup.title.get_text(' ', strip=True) if soup.title else path.stem),
                                     *[f'Embedded file: {url}' for url in embedded]])
            chunks = [('HTML body', content, '', '')]
        else:
            chunks = [('HTML text', soup.get_text('\n', strip=True), '', '')]
    elif suffix == '.pdf':
        from pypdf import PdfReader
        for index, page in enumerate(PdfReader(path).pages):
            text = page.extract_text() or ''
            if not text.strip():
                warnings.append(f'page {index+1}: no text layer; OCR needed')
            else:
                chunks.append((f'page {index+1}', text, '', ''))
    elif suffix == '.docx':
        from docx import Document
        document = Document(path)
        paragraphs = [p.text for p in document.paragraphs]
        paragraphs += [' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        chunks = [('document', '\n'.join(paragraphs), '', '')]
    elif suffix == '.xlsx':
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                rows = [' | '.join('' if v is None else str(v) for v in row) for row in sheet.iter_rows(values_only=True)]
                chunks.append((sheet.title, '\n'.join(rows), '', ''))
        finally:
            workbook.close()
    else:
        return [], [f'{suffix or "binary"}: parser not configured; original retained']
    # Bounded index rows preserve context without forcing one huge FTS row.
    bounded = []
    for heading, text, role, evidence in chunks:
        for start in range(0, len(text), 6000):
            if text[start:start+6000].strip():
                bounded.append((heading, text[start:start+6000], role, evidence))
    return bounded, warnings


def parse_version(root, version_id):
    base = core.root_path(root)
    con = core.connect(base, readonly=True)
    try:
        row = con.execute('SELECT v.*,d.evidence_kind FROM versions v JOIN documents d USING(document_id) WHERE version_id=?', (version_id,)).fetchone()
        if row is None:
            raise KeyError(version_id)
        row = dict(row)
        previous = con.execute('SELECT * FROM document_processing WHERE version_id=?', (version_id,)).fetchone()
        previous = dict(previous) if previous else None
        old_fragments = [dict(r) for r in con.execute('SELECT fragment_row_id,sequence_no FROM fragments WHERE version_id=?', (version_id,))]
    finally:
        con.close()
    path = base / row['raw_path']
    if any(f['fragment_row_id'] != f'{version_id}:part:{f["sequence_no"]}' for f in old_fragments):
        return {'status': previous['state'] if previous else 'existing_evidence', 'fragments': len(old_fragments),
                'detail': 'Preserved existing extracted evidence and anchors. Replacing a legacy extraction requires an explicit anchor mapping.'}
    integrity_verified = False
    try:
        if core.sha256(path) != row['content_sha256']:
            raise ValueError('Original checksum mismatch; extraction refused')
        integrity_verified = True
        chunks, warnings = extract(path)
        state = 'partial' if chunks and warnings else 'full_text_searchable' if chunks else 'pending'
        error = '; '.join(warnings) or None
    except Exception as exc:
        if previous and previous['state'].startswith('legacy_raw_') and not integrity_verified:
            return {'status': previous['state'], 'fragments': 0, 'detail': previous['error'],
                    'attempt_error': f'{type(exc).__name__}: {exc}', 'baseline_preserved': True}
        chunks, state, error = [], 'failed', f'{type(exc).__name__}: {exc}'
    con = core.connect(base)
    try:
        def write(c):
            check = c.execute('SELECT updated_at FROM document_processing WHERE version_id=?', (version_id,)).fetchone()
            if previous and check and check[0] != previous['updated_at']:
                raise core.ConflictError('Another extraction updated this version; re-read its state')
            current = c.execute('SELECT is_current FROM versions WHERE version_id=?', (version_id,)).fetchone()[0]
            # A failed retry must not erase previously successful extracted evidence.
            if chunks:
                expected_ids = {f'{version_id}:part:{i+1}' for i in range(len(chunks))}
                stale = [f['fragment_row_id'] for f in old_fragments if f['fragment_row_id'] not in expected_ids]
                for fragment_id in stale:
                    if c.execute('SELECT 1 FROM deadlines WHERE fragment_row_id=? LIMIT 1', (fragment_id,)).fetchone():
                        raise core.ConflictError('Re-extraction would remove referenced evidence; explicit mapping required')
                c.execute('DELETE FROM fragments_fts WHERE version_id=?', (version_id,))
                for fragment_id in stale:
                    c.execute('DELETE FROM fragments WHERE fragment_row_id=?', (fragment_id,))
                c.execute('DELETE FROM search_entries WHERE item_id=? AND version_id=? AND entry_key!=?', (row['document_id'], version_id, 'legacy:version:' + version_id))
                topics = [r[0] for r in c.execute('SELECT topic_id FROM document_topics WHERE document_id=?', (row['document_id'],))]
                for index, (heading, text, role, evidence) in enumerate(chunks):
                    fragment_id = f'{row["document_id"]}:part:{index+1}'
                    fragment_row_id = f'{version_id}:part:{index+1}'
                    c.execute('''INSERT INTO fragments(fragment_row_id,fragment_id,version_id,document_id,sequence_no,fragment_kind,heading,author_role,evidence_kind,content)
                                 VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(fragment_row_id) DO UPDATE SET
                                 heading=excluded.heading,author_role=excluded.author_role,evidence_kind=excluded.evidence_kind,content=excluded.content''',
                              (fragment_row_id, fragment_id, version_id, row['document_id'], index+1, 'text', heading, role, evidence or row['evidence_kind'], text))
                    title = c.execute('SELECT title FROM documents WHERE document_id=?', (row['document_id'],)).fetchone()[0]
                    c.execute('INSERT INTO fragments_fts(fragment_row_id,fragment_id,version_id,document_id,title,heading,content) VALUES(?,?,?,?,?,?,?)',
                              (fragment_row_id,fragment_id,version_id,row['document_id'],title,heading,text))
                    core.index_entry(c, row['document_id'], text, entry_key='legacy:fragment:' + fragment_row_id,
                                     heading=heading, source_ref=row['raw_path'] + '#' + heading, version_id=version_id,
                                     is_current=current, author_role=role, evidence_kind=evidence or row['evidence_kind'], topics=topics,
                                     authority='derived' if role == 'assistant' else 'user' if role == 'user' else None)
            c.execute('''INSERT INTO document_processing VALUES(?,?,?,?,?) ON CONFLICT(version_id) DO UPDATE SET
                         state=excluded.state,error=excluded.error,extractor=excluded.extractor,updated_at=excluded.updated_at''',
                      (version_id, state, error, 'v2-native-1', core.now()))
        core.transaction(con, write)
    finally:
        con.close()
    return {'status': state, 'fragments': len(chunks), 'detail': error}
